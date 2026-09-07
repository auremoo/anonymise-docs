"""Utilitaires partagés par la suite de tests.

Les documents de test (docx, pdf, images) sont **générés à l'exécution** :
le `.gitignore` exclut `*.docx` et `*.pdf` pour éviter de commiter des
documents sensibles par accident, donc aucun fixture binaire n'est stocké
dans le dépôt.
"""

import sys
from pathlib import Path

# Permet `import anonymize` quel que soit le répertoire d'exécution.
RACINE = Path(__file__).resolve().parent.parent
if str(RACINE) not in sys.path:
    sys.path.insert(0, str(RACINE))

try:
    import fitz  # PyMuPDF
    A_PDF = True
except ImportError:
    A_PDF = False

try:
    from docx import Document
    from docx.shared import Inches
    A_DOCX = True
except ImportError:
    A_DOCX = False


# Couleurs unies : permettent d'identifier chaque image par son pixel et
# donc de vérifier la correspondance placeholder <-> fichier extrait.
COULEURS = {
    "rouge": (255, 0, 0),
    "vert": (0, 255, 0),
    "bleu": (0, 0, 255),
}


def png_uni(nom_couleur: str, taille: int = 60) -> bytes:
    """PNG d'une couleur unie, identifiable par lecture d'un pixel."""
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, taille, taille))
    pix.set_rect(pix.irect, COULEURS[nom_couleur])
    return pix.tobytes("png")


def couleur_du_fichier(chemin: Path) -> str:
    """Nom de la couleur dominante d'une image (pour les assertions)."""
    pix = fitz.Pixmap(str(chemin))
    x = min(20, pix.width - 1)
    y = min(20, pix.height - 1)
    r, g, b = pix.pixel(x, y)[:3]
    if r > 200 and g < 60 and b < 60:
        return "rouge"
    if g > 200 and r < 60 and b < 60:
        return "vert"
    if b > 200 and r < 60 and g < 60:
        return "bleu"
    return f"inconnue({r},{g},{b})"


def creer_pdf(destination: Path) -> Path:
    """PDF 2 pages : texte + 1 image page 1, 2 images page 2.

    Ordre attendu des placeholders : rouge, vert, bleu.
    """
    doc = fitz.open()
    p1 = doc.new_page()
    p1.insert_text((50, 60), "Rapport pour Nexans, site de Bourg-en-Bresse.")
    p1.insert_text((50, 80), "Contact jean.dupont@acme.com au 06 12 34 56 78.")
    p1.insert_image(fitz.Rect(50, 100, 110, 160), stream=png_uni("rouge"))
    p2 = doc.new_page()
    p2.insert_text((50, 60), "Page 2 : architecture, IP 10.20.30.40.")
    p2.insert_image(fitz.Rect(50, 100, 110, 160), stream=png_uni("vert"))
    p2.insert_image(fitz.Rect(150, 100, 210, 160), stream=png_uni("bleu"))
    doc.save(str(destination))
    doc.close()
    return destination


def creer_docx(destination: Path, dossier_images: Path) -> Path:
    """docx couvrant les emplacements d'images qui ont posé problème.

    Ordre attendu des placeholders :
      IMAGE_1 = logo d'en-tête (rouge)
      IMAGE_2 = paragraphe (rouge)
      IMAGE_3 = cellule de tableau (vert)
      IMAGE_4 = paragraphe (bleu)
    """
    dossier_images.mkdir(parents=True, exist_ok=True)
    fichiers = {}
    for nom in ("rouge", "vert", "bleu"):
        f = dossier_images / f"src_{nom}.png"
        f.write_bytes(png_uni(nom))
        fichiers[nom] = str(f)

    doc = Document()
    doc.add_paragraph("Rapport pour Nexans, contact jean.dupont@acme.com.")
    doc.add_paragraph("Image dans un paragraphe :")
    doc.add_picture(fichiers["rouge"], width=Inches(1))

    doc.add_paragraph("Tableau contenant une image :")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Site"
    table.cell(0, 1).text = "Bourg-en-Bresse"
    table.cell(1, 0).text = "Schema"
    # Image DANS une cellule : jamais extraite avant correction.
    table.cell(1, 1).paragraphs[0].add_run().add_picture(
        fichiers["vert"], width=Inches(1)
    )

    doc.add_paragraph("Derniere image dans un paragraphe :")
    doc.add_picture(fichiers["bleu"], width=Inches(1))

    # Logo d'en-tête : ses relations d'images vivent dans une autre partie
    # du docx que le corps.
    entete = doc.sections[0].header
    entete.paragraphs[0].add_run().add_picture(
        fichiers["rouge"], width=Inches(0.5)
    )

    doc.save(str(destination))
    return destination
